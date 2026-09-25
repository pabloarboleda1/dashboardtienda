"""
Genera el Excel de reporte (mismo formato de siempre: Consignaciones,
Inventario, Retribución FONDO) a partir de los datos reales de la base
de datos, para un rango de fechas cualquiera.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

XL_COLOR_HEADER = "2F4F8F"
XL_COLOR_SUBTOTAL = "D9E1F2"
XL_FUENTE = "Arial"

CANAL_A_INVENTARIO = {
    "Física (Consignación)": "fisica",
    "Online (En línea)": "online",
    "Deducción Nómina": "nomina",
}
CANAL_A_MEDIO_PAGO = {
    "Física (Consignación)": "CONSIGNACIÓN",
    "Online (En línea)": "EN LINEA",
    "Deducción Nómina": "NÓMINA",
}


def _borde(cell):
    thin = Side(style="thin", color="BFBFBF")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def _hs(cell, color=XL_COLOR_HEADER):
    cell.font = Font(name=XL_FUENTE, bold=True, color="FFFFFF", size=10)
    cell.fill = PatternFill("solid", start_color=color, end_color=color)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _borde(cell)


def _ss(cell):
    cell.font = Font(name=XL_FUENTE, bold=True, size=10)
    cell.fill = PatternFill("solid", start_color=XL_COLOR_SUBTOTAL, end_color=XL_COLOR_SUBTOTAL)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    _borde(cell)


def _ds(cell, bg=None):
    cell.font = Font(name=XL_FUENTE, size=10)
    if bg:
        cell.fill = PatternFill("solid", start_color=bg, end_color=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _borde(cell)


def generar_excel_reporte(ventas, productos_activos):
    """ventas: lista de objetos Venta (con .producto ya cargado). productos_activos: lista de Producto."""
    ventas_por_producto = {}
    for v in ventas:
        ventas_por_producto.setdefault(v.producto_id, []).append(v)

    wb = Workbook()

    # ---------------- Hoja Consignaciones ----------------
    ws1 = wb.active
    ws1.title = "Consignaciones"
    encabezados1 = ["RETRIBUCIÓN", "MEDIO DE PAGO", "CLIENTE", "IDENTIFICACIÓN",
                     "VALOR MERCANCIA", "CONSIGNACIÓN", "FECHA DE PAGO", "DESCRIPCION"]
    for j, h in enumerate(encabezados1, start=1):
        _hs(ws1.cell(row=1, column=j, value=h))
    fila = 2
    for v in sorted(ventas, key=lambda x: x.fecha):
        p = v.producto
        retribucion_txt = "FONDO" if (p and (p.compra_origen or "").upper() == "FONDO DE EMPLEADOS") else "UCN"
        valores = [
            retribucion_txt, CANAL_A_MEDIO_PAGO.get(v.canal, v.canal or ""),
            v.cliente_nombre or "", v.cliente_documento or "",
            v.valor_venta or 0, v.valor_venta or 0, v.fecha,
            p.nombre_completo if p else "(producto eliminado)",
        ]
        for j, val in enumerate(valores, start=1):
            celda = ws1.cell(row=fila, column=j, value=val)
            if j in (5, 6):
                celda.number_format = "$#,##0"
            _ds(celda)
        fila += 1
    if not ventas:
        ws1.cell(row=2, column=1, value="Sin ventas registradas en este periodo")
    for col_l, w in zip("ABCDEFGH", [14, 16, 22, 16, 16, 16, 14, 32]):
        ws1.column_dimensions[col_l].width = w

    # ---------------- Hoja Inventario ----------------
    ws2 = wb.create_sheet("Inventario")
    encabezados2 = ["Producto", "Compra Origen", "Costo Unitario", "Precio Venta Unitario",
                     "Cantidad Vendida Física", "Cantidad Vendida Online", "Cant Deducción Nómina",
                     "Total Vendido", "Ingresos Totales", "Inventario Final (actual)", "Utilidad Bruta (del periodo)"]
    for j, h in enumerate(encabezados2, start=1):
        _hs(ws2.cell(row=1, column=j, value=h))
    fila = 2
    for p in sorted(productos_activos, key=lambda x: (x.compra_origen or "", x.nombre, x.talla)):
        ventas_p = ventas_por_producto.get(p.id, [])
        cant = {"fisica": 0, "online": 0, "nomina": 0}
        ingresos, costo_periodo = 0.0, 0.0
        for v in ventas_p:
            clave = CANAL_A_INVENTARIO.get(v.canal)
            if clave:
                cant[clave] += v.cantidad
            ingresos += v.valor_venta or 0
            costo_periodo += v.costo_total or 0
        total_vendido = sum(cant.values())
        nombre_export = p.nombre_completo + (" F" if (p.compra_origen or "").upper() == "FONDO DE EMPLEADOS" else "")

        valores = [
            nombre_export, p.compra_origen or "", p.costo_unitario or 0, p.precio_venta_unitario or 0,
            cant["fisica"], cant["online"], cant["nomina"], total_vendido, ingresos,
            p.stock_actual or 0, ingresos - costo_periodo,
        ]
        for j, val in enumerate(valores, start=1):
            celda = ws2.cell(row=fila, column=j, value=val)
            if j in (3, 4, 9, 11):
                celda.number_format = "$#,##0"
            _ds(celda)
        fila += 1
    for col_l, w in zip("ABCDEFGHIJK", [30, 20, 14, 16, 12, 12, 12, 12, 16, 18, 20]):
        ws2.column_dimensions[col_l].width = w

    fila += 1
    nota = ws2.cell(row=fila, column=1, value=(
        "Nota: 'Inventario Final' muestra el stock ACTUAL (al momento de exportar), no el que había "
        "exactamente al cierre de este periodo — eso requiere el historial de Compras, que aún no está "
        "implementado. 'Utilidad Bruta' sí es exacta: usa el costo real de cada venta del periodo."
    ))
    nota.font = Font(name=XL_FUENTE, italic=True, size=9, color="7F7F7F")
    nota.alignment = Alignment(wrap_text=True)
    ws2.merge_cells(f"A{fila}:K{fila}")
    ws2.row_dimensions[fila].height = 30

    # ---------------- Hoja Retribución FONDO ----------------
    ws3 = wb.create_sheet("Retribución FONDO")
    encabezados3 = ["Producto", "Categoría", "Unidades Vendidas", "Retribución"]
    for j, h in enumerate(encabezados3, start=1):
        _hs(ws3.cell(row=1, column=j, value=h))
    fila = 2
    total_retribucion = 0.0
    for p in productos_activos:
        ventas_p = [v for v in ventas_por_producto.get(p.id, []) if v.retribucion_fondo]
        if not ventas_p:
            continue
        unidades = sum(v.cantidad for v in ventas_p)
        retribucion = sum(v.retribucion_fondo for v in ventas_p)
        valores = [p.nombre_completo, p.categoria or "", unidades, retribucion]
        for j, val in enumerate(valores, start=1):
            celda = ws3.cell(row=fila, column=j, value=val)
            if j == 4:
                celda.number_format = "$#,##0"
            _ds(celda)
        total_retribucion += retribucion
        fila += 1

    if fila == 2:
        ws3.cell(row=2, column=1, value="No hay productos del Fondo con ventas en este periodo")
    else:
        ws3.merge_cells(f"A{fila}:C{fila}")
        _ss(ws3.cell(row=fila, column=1, value="TOTAL A RETRIBUIR"))
        c_total = ws3.cell(row=fila, column=4, value=total_retribucion)
        _ss(c_total)
        c_total.number_format = "$#,##0"
    for col_l, w in zip("ABCD", [32, 22, 16, 16]):
        ws3.column_dimensions[col_l].width = w

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    return salida
